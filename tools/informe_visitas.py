#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Convierte la exportación de GoatCounter (el .zip que se descarga en Settings → Export)
en un Excel con los resúmenes para la memoria del proyecto: visitas por página, por día,
por mes, por dispositivo, por sistema, por navegador, por país, por origen y clics en botones.

Uso:  python tools/informe_visitas.py ruta/goatcounter-export-....zip [salida.xlsx]

El .zip lleva el correo de la cuenta: no se guarda en el repositorio.
"""
import collections
import io
import json
import sys
import zipfile
from pathlib import Path
from urllib.parse import unquote

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def leer(z, nombre):
    ruta = next((n for n in z.namelist() if n.endswith('/' + nombre) or n == nombre), None)
    if not ruta:
        return []
    return [json.loads(l) for l in io.TextIOWrapper(z.open(ruta), encoding='utf-8') if l.strip()]


def dispositivo(ancho):
    """Tamaño de pantalla (ancho en píxeles) → tipo de dispositivo."""
    if not ancho:
        return 'Desconocido'
    if ancho < 768:
        return 'Móvil'
    if ancho < 1280:
        return 'Tableta'
    return 'Ordenador'


def nombre_pagina(p):
    if p.get('event'):
        return p['path']
    ruta = unquote(p['path']).replace('/biocelia/', '', 1).replace('/biocelia', 'index.html')
    return ruta or 'index.html'


def hoja(libro, titulo, cabecera, filas, nota=None):
    h = libro.create_sheet(titulo)
    fila0 = 1
    if nota:
        h.cell(1, 1, nota).font = Font(italic=True, color='555555')
        fila0 = 3
    for j, c in enumerate(cabecera, 1):
        celda = h.cell(fila0, j, c)
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='0AA3C2')
    for i, f in enumerate(filas, fila0 + 1):
        for j, v in enumerate(f, 1):
            h.cell(i, j, v)
    for j in range(1, len(cabecera) + 1):
        ancho = max(len(str(h.cell(i, j).value or '')) for i in range(fila0, fila0 + len(filas) + 1))
        h.column_dimensions[get_column_letter(j)].width = min(max(12, ancho + 2), 70)
    h.freeze_panes = h.cell(fila0 + 1, 1)
    return h


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    origen = Path(sys.argv[1])
    salida = Path(sys.argv[2]) if len(sys.argv) > 2 else origen.with_name('informe-visitas-' + origen.stem.split('-')[-1][:8] + '.xlsx')
    z = zipfile.ZipFile(origen)

    paginas = {p['id']: p for p in leer(z, 'paths.jsonl')}
    refs = {r['id']: r['ref'] for r in leer(z, 'refs.jsonl')}
    sistemas = {s['id']: s['name'] for s in leer(z, 'systems.jsonl')}
    navegadores = {b['id']: b['name'] for b in leer(z, 'browsers.jsonl')}
    paises = {}
    for l in leer(z, 'locations.jsonl'):
        paises.setdefault(l['country'], l.get('country_name') or l['country'])

    visitas = leer(z, 'hit_stats.jsonl')
    es_evento = lambda pid: bool(paginas.get(pid, {}).get('event'))
    vistas = [v for v in visitas if not es_evento(v['path_id'])]
    clics = [v for v in visitas if es_evento(v['path_id'])]

    libro = Workbook()
    libro.remove(libro.active)

    total = sum(v['count'] for v in vistas)
    dias = sorted({v['hour'][:10] for v in visitas})
    hoja(libro, 'Resumen', ['Dato', 'Valor'], [
        ['Periodo', f"{dias[0]} a {dias[-1]}" if dias else '—'],
        ['Visitas a páginas', total],
        ['Clics registrados en botones', sum(v['count'] for v in clics)],
        ['Páginas distintas visitadas', len({v['path_id'] for v in vistas})],
        ['Días con visitas', len({v['hour'][:10] for v in vistas})],
    ], 'Resumen de la exportación de GoatCounter (visitas anónimas, sin cookies ni datos personales).')

    # «/biocelia» y «/biocelia/index.html» son la misma portada: se cuentan juntas
    c, titulos = collections.Counter(), {}
    for v in vistas:
        nombre = nombre_pagina(paginas[v['path_id']])
        c[nombre] += v['count']
        titulos.setdefault(nombre, paginas[v['path_id']].get('title', ''))
    hoja(libro, 'Por página', ['Página', 'Título', 'Visitas', '% del total'],
         [[k, titulos[k], n, round(100 * n / total, 1) if total else 0] for k, n in c.most_common()])

    c = collections.Counter()
    for v in vistas:
        c[v['hour'][:10]] += v['count']
    hoja(libro, 'Por día', ['Día', 'Visitas'], sorted(c.items()))

    c = collections.Counter()
    for v in vistas:
        c[v['hour'][:7]] += v['count']
    hoja(libro, 'Por mes', ['Mes', 'Visitas'], sorted(c.items()))

    def por(stats, clave, nombre):
        cc = collections.Counter()
        for s in leer(z, stats):
            if not es_evento(s['path_id']):
                cc[nombre(s.get(clave))] += s['count']
        return cc

    c = por('size_stats.jsonl', 'width', dispositivo)
    t = sum(c.values()) or 1
    hoja(libro, 'Por dispositivo', ['Dispositivo', 'Visitas', '%'],
         [[k, n, round(100 * n / t, 1)] for k, n in c.most_common()],
         'Según el ancho de pantalla: móvil < 768 px, tableta 768-1279 px, ordenador ≥ 1280 px.')
    c = por('system_stats.jsonl', 'system_id', lambda i: sistemas.get(i, 'Desconocido'))
    hoja(libro, 'Por sistema', ['Sistema', 'Visitas'], c.most_common())
    c = por('browser_stats.jsonl', 'browser_id', lambda i: navegadores.get(i, 'Desconocido'))
    hoja(libro, 'Por navegador', ['Navegador', 'Visitas'], c.most_common())
    c = por('location_stats.jsonl', 'location', lambda l: paises.get(l, l) if l else 'Desconocido')
    hoja(libro, 'Por país', ['País', 'Visitas'], c.most_common())

    c = collections.Counter()
    for v in vistas:
        c[refs.get(v.get('ref_id'), '') or '(entrada directa o sin origen)'] += v['count']
    hoja(libro, 'Por origen', ['Desde dónde llegaron', 'Visitas'], c.most_common())

    c = collections.Counter()
    for v in clics:
        c[paginas[v['path_id']]['path']] += v['count']
    hoja(libro, 'Clics en botones', ['Botón', 'Clics'], c.most_common(),
         'Pdf-…: descargas de PDF; juego-…: juego abierto; portada-…: tarjetas de la portada.')

    for h in libro.worksheets:
        for fila in h.iter_rows():
            for celda in fila:
                celda.alignment = Alignment(vertical='top', wrap_text=False)
    libro.save(salida)
    print('Informe guardado en', salida)
    return 0


if __name__ == '__main__':
    sys.exit(main())
